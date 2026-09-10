"""Build the PRIMARY demo fixture pack — the plant floor story.

Deliberately plain: machines, temperature, a policy that says what "too hot"
means. A juror with no domain background should understand the answer in one
glance, which the corrosion/ultrasonic pack does not deliver.

Everything is invented. Every page is stamped SYNTHETIC.

Run inside the worker container:

    docker compose run --rm --no-deps \
      -v "D:/Code/psu/fixtures:/app/fixtures" worker \
      python fixtures/generate_demo_fixtures.py

Outputs:
    machine_health_log.xlsx    12 machines; 3 are over the temperature limit
    maintenance_policy.pdf     plain-English policy with the 85 degrees C rule
    demo_answers.json          what a correct run must produce
"""
from __future__ import annotations

import json
from pathlib import Path

import pymupdf

FIXTURES = Path(__file__).resolve().parent
STAMP = "SYNTHETIC DEMO DATA - NOT A REAL PLANT RECORD"

# machine_id, name, temperature_c, vibration_mm_s, units_produced, downtime_minutes
MACHINES = [
    ("M-01", "Injection Moulder 1",  72.4, 2.1, 1840,  0),
    ("M-02", "Injection Moulder 2",  78.9, 2.8, 1795, 15),
    ("M-03", "Conveyor Drive A",     64.2, 1.4, 2010,  0),
    ("M-04", "Hydraulic Press 1",    91.6, 5.2, 1120, 95),
    ("M-05", "Hydraulic Press 2",    83.1, 3.9, 1680, 20),
    ("M-06", "Packaging Line 1",     69.7, 2.0, 2240,  0),
    ("M-07", "Air Compressor",       88.3, 4.7,    0, 45),
    ("M-08", "Conveyor Drive B",     66.5, 1.6, 1990,  0),
    ("M-09", "Cooling Pump 1",       74.8, 2.4, 1570, 10),
    ("M-10", "Grinding Unit",        94.2, 6.1,  880, 130),
    ("M-11", "Packaging Line 2",     71.0, 2.2, 2180,  0),
    ("M-12", "Cooling Pump 2",       76.3, 2.6, 1610,  5),
]

TEMPERATURE_LIMIT = 85.0
VIBRATION_LIMIT = 4.5

POLICY = f"""{STAMP}

PLANT MAINTENANCE POLICY
Document MP-2026-01  |  Revision 3  |  Effective 1 January 2026

1. PURPOSE

1.1 This policy defines when a machine on the production floor must be taken
out of service for maintenance, and who is allowed to authorise that.

2. TEMPERATURE LIMITS

2.1 Normal operating temperature for all production machinery is 85 degrees
Celsius or below.

2.2 A machine recorded above 85 degrees Celsius is classified as OVERHEATING
and must be scheduled for maintenance within 24 hours.

2.3 A machine recorded above 90 degrees Celsius must be stopped immediately
and may not be restarted until an engineer has inspected it.

3. VIBRATION LIMITS

3.1 Vibration above 4.5 mm/s indicates bearing wear and requires inspection at
the next scheduled stop.

3.2 A machine exceeding both the temperature and the vibration limit is treated
as HIGH PRIORITY and is inspected before any other machine.

4. AUTHORISATION

4.1 Any maintenance shutdown must be approved by the shift supervisor before
the machine is stopped.

4.2 No automated system, dashboard or reporting tool may record that approval
on the supervisor's behalf. The supervisor signs personally.

4.3 The signed approval note is filed against the machine record and kept for
two years.

5. RECORDS

5.1 Daily machine readings are logged at the end of every shift.

5.2 Machines flagged OVERHEATING are re-checked on the following shift and the
result is added to the same record.
"""

OVER_TEMPERATURE = [m for m in MACHINES if m[2] > TEMPERATURE_LIMIT]
OVER_BOTH = [m for m in OVER_TEMPERATURE if m[3] > VIBRATION_LIMIT]

DEMO_ANSWERS = {
    "_note": "What a correct run must produce. All data is synthetic.",
    "primary_question": (
        "Which machines are running above 85 degrees C? Check them against our "
        "maintenance policy and draft a maintenance approval note."
    ),
    "expected_route": {"task_class": "spreadsheet-analysis", "model_id": "qwen3-8b"},
    "temperature_limit_c": TEMPERATURE_LIMIT,
    "expected_match_count": len(OVER_TEMPERATURE),
    "expected_machines": [m[0] for m in OVER_TEMPERATURE],
    "expected_machine_names": {m[0]: m[1] for m in OVER_TEMPERATURE},
    "expected_temperatures": {m[0]: m[2] for m in OVER_TEMPERATURE},
    "high_priority_both_limits": [m[0] for m in OVER_BOTH],
    "expected_policy_clauses": ["2.2", "2.3", "4.1", "4.2"],
    "expected_citation_count_min": 2,
    "must_pause_for_human_approval": True,
    "deliverable": "maintenance_note.docx",
}


def build_policy(out_path: Path) -> None:
    # Two pages, so citations carry more than one page number.
    head, tail = POLICY.split("4. AUTHORISATION", 1)
    pages = [
        head.rstrip(),
        f"{STAMP}\n\nPLANT MAINTENANCE POLICY (continued)\n\n4. AUTHORISATION{tail}",
    ]

    doc = pymupdf.open()
    for body in pages:
        page = doc.new_page(width=595, height=842)  # A4
        page.insert_textbox(
            pymupdf.Rect(56, 56, 539, 786),
            body,
            fontname="cour",
            fontsize=10.5,
            align=pymupdf.TEXT_ALIGN_LEFT,
        )
    doc.save(str(out_path))
    print(f"  {out_path.name}: {doc.page_count} pages, plain-English policy")
    doc.close()


def build_machine_log(out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Daily Readings"

    headers = [
        "machine_id", "machine_name", "temperature_c", "vibration_mm_s",
        "units_produced", "downtime_minutes", "shift",
    ]
    sheet.append(headers)
    for machine in MACHINES:
        sheet.append([*machine, "Shift A"])

    summary = wb.create_sheet("Shift Summary")
    summary.append(["metric", "value"])
    for row in [
        ("date", "2026-09-10"),
        ("shift", "A"),
        ("machines_logged", len(MACHINES)),
        ("total_units_produced", sum(m[4] for m in MACHINES)),
        ("total_downtime_minutes", sum(m[5] for m in MACHINES)),
        ("temperature_limit_c", TEMPERATURE_LIMIT),
        ("vibration_limit_mm_s", VIBRATION_LIMIT),
        ("logged_by", "Floor Supervisor (synthetic)"),
    ]:
        summary.append(row)

    header_fill = PatternFill("solid", fgColor="1E2130")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="F8FAFC")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left")
        ws.freeze_panes = "A2"
        for column_cells in ws.columns:
            width = max(len(str(c.value or "")) for c in column_cells) + 3
            ws.column_dimensions[column_cells[0].column_letter].width = min(width, 28)

    wb.save(out_path)
    print(f"  {out_path.name}: {len(MACHINES)} machines, "
          f"{len(OVER_TEMPERATURE)} above {TEMPERATURE_LIMIT} C "
          f"({', '.join(m[0] for m in OVER_TEMPERATURE)})")


def main() -> None:
    print(f"Writing demo fixtures into {FIXTURES}")
    build_machine_log(FIXTURES / "machine_health_log.xlsx")
    build_policy(FIXTURES / "maintenance_policy.pdf")
    (FIXTURES / "demo_answers.json").write_text(
        json.dumps(DEMO_ANSWERS, indent=2), encoding="utf-8"
    )
    print("  demo_answers.json: expected results")
    print("Done.")


if __name__ == "__main__":
    main()
