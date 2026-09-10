"""Approval-note DOCX generation, verification, and post-decision stamping."""
from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List

from langchain_core.tools import tool
from pydantic import BaseModel, Field

from worker.tools.path_utils import PathPolicyError, relative_to_workspace, validate_path

MIN_DOCX_BYTES = 5 * 1024
MIN_CITATIONS = 2

MANDATORY_HEADINGS = [
    "Purpose",
    "Source Report Metadata",
    "Extracted Findings",
    "Applicable SOP Clauses",
    "Recommended Action",
    "Risk Notice",
    "Approval",
]

RISK_NOTICE = (
    "This note was drafted by an automated decision-support agent. It is NOT an "
    "engineering certification and carries no authority on its own. It requires "
    "authorized engineer approval before any physical action is taken."
)

# "(Source: corrosion_sop.pdf, Page 14)"
CITATION_RE = re.compile(r"\(Source:\s*[^,]+,\s*Page\s*[^)]+\)", re.IGNORECASE)


class Finding(BaseModel):
    item: str = Field(..., description="Component or item the finding concerns")
    detail: str = Field(..., description="What was observed, in the source's own terms")
    severity: str = Field("Info", description="High / Medium / Low / Info")
    evidence: str = Field("", description="Where this came from, e.g. 'OCR page 2'")


class SopClause(BaseModel):
    clause: str = Field(..., description="Clause identifier, e.g. SOP-802-A ss4.2")
    content: str = Field(..., description="The clause text as retrieved")
    source_file: str = Field(..., description="Source document filename")
    page_number: str | int = Field(..., description="Page number in the source document")


class CreateDocxInput(BaseModel):
    output_filename: str = Field(..., description="e.g. approval_note.docx")
    purpose: str = Field(..., description="Why this note exists, one short paragraph")
    source_metadata: Dict[str, Any] = Field(..., description="Source report metadata")
    findings: List[Finding] = Field(..., description="Findings extracted from the source")
    sop_clauses: List[SopClause] = Field(..., description="Retrieved SOP clauses, with citations")
    recommended_action: str = Field(..., description="What the agent recommends a human do")
    run_id: str = Field(..., description="Run ID")


class VerifyDocxInput(BaseModel):
    docx_path: str = Field(..., description="Workspace-relative path to the DOCX")
    run_id: str = Field(..., description="Run ID")


def _artifact_path(output_filename: str, run_id: str) -> Path:
    name = Path(output_filename).name  # never let a directory ride along
    if not name.lower().endswith(".docx"):
        name += ".docx"
    return validate_path(Path("artifacts") / name, run_id)


@tool("create_approval_docx", args_schema=CreateDocxInput)
def create_approval_docx(
    output_filename: str,
    purpose: str,
    source_metadata: Dict[str, Any],
    findings: List[Finding],
    sop_clauses: List[SopClause],
    recommended_action: str,
    run_id: str,
) -> Dict[str, Any]:
    """Generate the formal approval note DOCX with every mandatory section and
    page-level citations for each SOP clause."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    try:
        out = _artifact_path(output_filename, run_id)
    except PathPolicyError as e:
        return {"status": "rejected", "error": str(e)}
    out.parent.mkdir(parents=True, exist_ok=True)

    # Pydantic gives models; a raw dict may arrive from a hand-built plan.
    findings = [f if isinstance(f, Finding) else Finding(**f) for f in findings]
    sop_clauses = [c if isinstance(c, SopClause) else SopClause(**c) for c in sop_clauses]

    doc = docx.Document()

    title = doc.add_heading("TECHNICAL APPROVAL NOTE (DRAFT)", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run(
        f"Generated locally by the Sovereign AI Workbench · Run {run_id} · "
        f"{datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    sub_run.italic = True
    sub_run.font.size = Pt(9)

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(purpose)

    doc.add_heading("Source Report Metadata", level=1)
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.style = "Table Grid"
    for key, value in source_metadata.items():
        cells = meta_table.add_row().cells
        cells[0].text = str(key).replace("_", " ").title()
        cells[1].text = str(value)

    doc.add_heading("Extracted Findings Table", level=1)
    if findings:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        headers = ["Item / Component", "Finding", "Severity", "Evidence"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            for para in cell.paragraphs:
                for r in para.runs:
                    r.bold = True
        for finding in findings:
            cells = table.add_row().cells
            cells[0].text = finding.item
            cells[1].text = finding.detail
            cells[2].text = finding.severity
            cells[3].text = finding.evidence or "—"
    else:
        doc.add_paragraph("No findings were extracted from the source document.")

    doc.add_heading("Applicable SOP Clauses", level=1)
    if sop_clauses:
        for clause in sop_clauses:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"[{clause.clause}] ").bold = True
            p.add_run(clause.content.strip() + " ")
            cite = p.add_run(f"(Source: {clause.source_file}, Page {clause.page_number})")
            cite.italic = True
    else:
        doc.add_paragraph(
            "No SOP clauses were retrieved from the local knowledge base. "
            "This note cannot be relied upon without them."
        )

    doc.add_heading("Recommended Action", level=1)
    doc.add_paragraph(recommended_action)

    doc.add_heading("Risk Notice", level=1)
    risk = doc.add_paragraph().add_run(RISK_NOTICE)
    risk.bold = True
    risk.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)

    doc.add_heading("Approval / Signature", level=1)
    doc.add_paragraph(
        "Status: PENDING — no human decision has been recorded for this run yet."
    )
    sig = doc.add_table(rows=2, cols=2)
    sig.style = "Table Grid"
    sig.rows[0].cells[0].text = "Authorized Engineer (name / signature):"
    sig.rows[0].cells[1].text = ""
    sig.rows[1].cells[0].text = "Date:"
    sig.rows[1].cells[1].text = ""

    doc.save(str(out))

    return {
        "status": "success",
        "docx_path": relative_to_workspace(out, run_id),
        "absolute_path": str(out),
        "file_size": out.stat().st_size,
        "findings_count": len(findings),
        "citation_count": len(sop_clauses),
    }


@tool("verify_docx", args_schema=VerifyDocxInput)
def verify_docx(docx_path: str, run_id: str) -> Dict[str, Any]:
    """Reopen the generated DOCX and check it is parseable and contains every
    mandatory heading plus at least two page-level source citations."""
    import docx

    try:
        path = validate_path(docx_path, run_id)
    except PathPolicyError as e:
        return {"valid": False, "reason": str(e)}

    if not path.exists():
        # Tolerate a bare filename by looking in the run's artifacts directory.
        alt = validate_path(Path("artifacts") / Path(docx_path).name, run_id)
        if alt.exists():
            path = alt
        else:
            return {"valid": False, "reason": f"File does not exist: {docx_path}"}

    size = path.stat().st_size
    if size < MIN_DOCX_BYTES:
        return {"valid": False, "reason": f"File is only {size} bytes (< {MIN_DOCX_BYTES})"}

    try:
        doc = docx.Document(str(path))
    except Exception as e:
        return {"valid": False, "reason": f"Not a readable DOCX: {e}"}

    headings = [
        p.text.strip() for p in doc.paragraphs
        if p.style is not None and p.style.name.startswith(("Heading", "Title"))
    ]
    missing = [
        h for h in MANDATORY_HEADINGS
        if not any(h.lower() in found.lower() for found in headings)
    ]

    body_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            body_text += "\n" + " ".join(c.text for c in row.cells)

    citations = CITATION_RE.findall(body_text)

    problems = []
    if missing:
        problems.append(f"missing mandatory headings: {missing}")
    if len(citations) < MIN_CITATIONS:
        problems.append(f"only {len(citations)} source citation(s), need {MIN_CITATIONS}")

    return {
        "valid": not problems,
        "reason": "; ".join(problems) if problems else "",
        "docx_path": relative_to_workspace(path, run_id),
        "absolute_path": str(path),
        "file_size": size,
        "headings_found": headings,
        "missing_headings": missing,
        "citation_count": len(citations),
        "citations": citations[:10],
    }


def stamp_human_decision(
    docx_path: str | Path,
    decision: str,
    note: str,
    decided_by: str = "human",
) -> Dict[str, Any]:
    """Write the *recorded* human decision into the note's Approval section.

    Deliberately NOT a LangChain tool: the agent must have no code path that can
    record or fabricate an approval. Only the DELIVER node calls this, and only
    after the decision exists in the `approval_decisions` table.
    """
    import docx

    path = Path(docx_path)
    if not path.exists():
        return {"status": "error", "error": f"DOCX not found: {path}"}

    document = docx.Document(str(path))
    decided_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    for paragraph in document.paragraphs:
        if paragraph.text.startswith("Status: PENDING"):
            for run in list(paragraph.runs):
                run.text = ""
            stamped = paragraph.add_run(
                f"Status: {decision.upper()} — recorded {decided_at} by {decided_by}."
                + (f" Note: {note}" if note else "")
            )
            stamped.bold = True
            break

    for table in document.tables:
        for row in table.rows:
            if row.cells[0].text.startswith("Authorized Engineer"):
                row.cells[1].text = f"{decided_by} ({decision.upper()})"
            elif row.cells[0].text.startswith("Date:"):
                row.cells[1].text = decided_at

    document.save(str(path))
    return {"status": "success", "decision": decision, "decided_at": decided_at,
            "docx_path": str(path)}
